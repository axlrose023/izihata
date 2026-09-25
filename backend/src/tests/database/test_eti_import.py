from pathlib import Path

import pytest
import pytest_asyncio
from openpyxl import Workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import selectinload

from app.api.modules.catalog.enums import ProductAttributeSource, StockStatus
from app.api.modules.catalog.models import Product
from app.database.imports import (
    BunnyS3Config,
    import_eti_workbook,
    planned_media_urls,
    read_eti_workbook,
)


def write_workbook(path: Path) -> Path:
    workbook = Workbook()
    products = workbook.active
    products.title = "products"
    products.append(["Код постачальника", "Назва товару", "Ціна в грн. з ПДВ"])
    products.append(["ETI-TEST-1", "Авт. вимикач ETIMAT 6 1p C16", 120.5])
    products.append(["ETI-TEST-1", "Авт. вимикач ETIMAT 6 1p C16", 120.5])

    characteristics = workbook.create_sheet("characteristics")
    characteristics.append(
        [
            "Код постачальника",
            "Джерело",
            "Характеристика",
            "Значення",
            "Число",
            "Одиниця",
        ]
    )
    characteristics.append(["ETI-TEST-1", "основні", "Полюси", "1P", None, None])
    characteristics.append(["ETI-TEST-1", "основні", "Номінал", None, 16, "А"])
    characteristics.append(["ETI-TEST-1", "ETIM", "Ширина", None, 18, "мм"])

    photos = workbook.create_sheet("photos")
    photos.append(["Код постачальника", "URL"])
    photos.append(["ETI-TEST-1", "https://supplier.example/one.webp"])
    photos.append(["ETI-TEST-1", "https://supplier.example/two.webp"])
    workbook.save(path)
    return path


@pytest.mark.asyncio
class TestEtiWorkbookImport:
    @pytest_asyncio.fixture(autouse=True)
    async def _clean_slate(self, uow):
        await uow.session.execute(delete(Product).where(Product.sku == "ETI-TEST-1"))
        await uow.commit()
        yield
        await uow.session.execute(delete(Product).where(Product.sku == "ETI-TEST-1"))
        await uow.commit()

    async def test_imports_primary_and_etim_specs_with_all_media(self, tmp_path, uow):
        workbook = read_eti_workbook(write_workbook(tmp_path / "eti.xlsx"))
        config = BunnyS3Config(
            endpoint="https://de-s3.storage.bunnycdn.com",
            storage_zone="izihata-product-media",
            password="test-password",
            public_base_url="https://izihata-product-media.b-cdn.net",
        )

        outcome = await import_eti_workbook(
            uow.session,
            workbook,
            media_urls=planned_media_urls(workbook, config),
            dry_run=False,
            batch_size=1,
        )

        product = (
            await uow.session.execute(
                select(Product)
                .where(Product.sku == "ETI-TEST-1")
                .options(selectinload(Product.attributes), selectinload(Product.media))
            )
        ).scalar_one()
        primary = [
            attribute
            for attribute in product.attributes
            if attribute.source == ProductAttributeSource.PRIMARY
        ]
        etim = [
            attribute
            for attribute in product.attributes
            if attribute.source == ProductAttributeSource.ETIM
        ]

        assert (outcome.created, outcome.updated) == (1, 0)
        assert workbook.duplicate_product_rows == 1
        assert product.brand == "ETI"
        assert product.stock_status == StockStatus.PREORDER
        assert product.is_active is True
        assert [(item.key, item.value) for item in primary] == [
            ("Код виробника", "ETI-TEST-1"),
            ("Виробник", "ETI"),
            ("Полюси", "1P"),
            ("Номінал", "16 А"),
        ]
        assert [(item.key, item.value) for item in etim] == [("Ширина", "18 мм")]
        assert [media.url for media in product.media] == [
            "https://izihata-product-media.b-cdn.net/products/eti/ETI-TEST-1/1.webp",
            "https://izihata-product-media.b-cdn.net/products/eti/ETI-TEST-1/2.webp",
        ]
        assert product.image_url == product.media[0].url

        repeated = await import_eti_workbook(
            uow.session,
            workbook,
            media_urls=planned_media_urls(workbook, config),
            dry_run=False,
            batch_size=1,
        )
        assert (repeated.created, repeated.updated) == (0, 1)

    async def test_dry_run_does_not_write_products(self, tmp_path, uow):
        workbook = read_eti_workbook(write_workbook(tmp_path / "eti.xlsx"))

        outcome = await import_eti_workbook(uow.session, workbook, dry_run=True)

        stored = (
            await uow.session.execute(
                select(Product).where(Product.sku == "ETI-TEST-1")
            )
        ).scalar_one_or_none()
        assert outcome.created == 1
        assert stored is None
